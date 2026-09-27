from __future__ import annotations

import typing
from typing import Any

from openpyxl import Workbook, load_workbook

from ..base import DEFAULT_BULK_NUMBER, BaseCodec, BaseFileIterable

_BLANK_HEADER_TOKENS = frozenset({"", "none", "null", "<na>", "nan"})

# Cell value types openpyxl accepts natively; anything else is stringified
# because XLSX rows are flat (``is_flatonly``).
_SCALAR_WRITE_TYPES = (str, int, float, bool, type(None))


def _is_blank_header(value: object) -> bool:
    """Return True when a header cell is missing or a non-name placeholder."""

    if value is None:
        return True
    text = str(value).strip()
    return not text or text.casefold() in _BLANK_HEADER_TOKENS


def read_row_keys(rownum, ncols, sheet):
    """Read single row by row num"""
    tmp = list()
    for i in range(0, ncols):
        sheet.cell_type(rownum, i)
        cell_value = sheet.cell_value(rownum, i)
        get_col = str(cell_value)
        tmp.append(get_col)
    return tmp


class XLSXIterable(BaseFileIterable):
    datamode = "binary"

    def __init__(
        self,
        filename: str | None = None,
        stream: typing.IO[Any] | None = None,
        codec: BaseCodec | None = None,
        mode="r",
        keys: list[str] | None = None,
        page: int = 0,
        start_line: int = 0,
        options: dict[str, Any] | None = None,
    ):
        if options is None:
            options = {}
        super().__init__(filename, stream, codec=codec, mode=mode, binary=True, noopen=True, options=options)
        self.keys = keys
        self.start_line = start_line + 1
        self.page = page
        if "page" in options.keys():
            self.page = int(options["page"])
        self.pos = self.start_line
        self.extracted_keys = keys is None
        self.reset()
        pass

    def reset(self):
        super().reset()
        if self.mode in ("w", "wr"):
            # Buffered writer: rows accumulate until close() materializes the
            # workbook (spreadsheet documents are single-file containers).
            self.workbook = None
            self.sheet = None
            self._write_rows: list[dict] = []
            self._write_keys: list[str] = list(self.keys) if self.keys else []
            return
        # read_only avoids loading the whole workbook into memory.
        self.workbook = load_workbook(self.filename, read_only=True)
        self.sheet = self.workbook[self.workbook.sheetnames[self.page]]
        self._total_rows = self.sheet.max_row
        self.pos = self.start_line
        # Some workbooks advertise incorrect dimensions (e.g. ref="A1").
        # Unbounded iter_rows() then returns a single cell; reset so openpyxl
        # discovers the real used range while streaming.
        reset_dimensions = getattr(self.sheet, "reset_dimensions", None)
        if callable(reset_dimensions):
            reset_dimensions()
        self.cursor = self.sheet.iter_rows()
        if self.pos > 1:
            self.skip(self.pos - 1)
        if self.extracted_keys:
            self.keys = list()
            try:
                # Skip leading blank rows (common before metadata tables), then
                # take the first non-blank row as headers. Sheets that never
                # produce a real header stay empty and are skipped upstream.
                blank_skips = 0
                while True:
                    row = next(self.cursor)
                    self.pos += 1
                    raw_keys = [cell.value for cell in row]
                    if not raw_keys or all(_is_blank_header(value) for value in raw_keys):
                        blank_skips += 1
                        if blank_skips >= 100:
                            self.keys = []
                            break
                        continue
                    self.keys = [str(value) for value in raw_keys]
                    break
            except StopIteration:
                # Empty sheet, use empty keys
                self.keys = []

    def skip(self, num: int = 1):
        while num > 0:
            num -= 1
            next(self.cursor)

    @staticmethod
    def has_totals() -> bool:
        """Has totals indicator"""
        return True

    def totals(self) -> int:
        """Returns file totals"""
        if getattr(self, "sheet", None) is not None and self.sheet.max_row is not None:
            return self.sheet.max_row
        total = getattr(self, "_total_rows", None)
        if total is not None:
            return total
        return 0

    @staticmethod
    def id() -> str:
        return "xlsx"

    @staticmethod
    def is_flatonly() -> bool:
        return True

    @staticmethod
    def has_tables() -> bool:
        """Indicates if this format supports multiple tables/sheets."""
        return True

    def list_tables(self, filename: str | None = None) -> list[str] | None:
        """List available sheet names in the Excel file.

        Can be called as:
        - Instance method: `iterable.list_tables()` - reuses open workbook
        - With filename: `iterable.list_tables(filename)` - opens file temporarily

        Args:
            filename: Optional filename. If None, uses instance's filename and reuses open workbook.

        Returns:
            list[str]: List of sheet names, or empty list if no sheets.
        """
        # If workbook is already open, reuse it
        if filename is None and hasattr(self, "workbook") and self.workbook is not None:
            return self.workbook.sheetnames

        # Otherwise, open temporarily
        target_filename = filename if filename is not None else self.filename
        if target_filename is None:
            return None

        workbook = load_workbook(target_filename, read_only=True)
        try:
            return workbook.sheetnames
        finally:
            workbook.close()

    def read(self, skip_empty: bool = True) -> dict:
        """Read single XLSX record"""
        row = next(self.cursor)
        tmp = list()
        for cell in row:
            tmp.append(str(cell.value))
        result = dict(zip(self.keys, tmp, strict=False))
        self.pos += 1
        return result

    def read_bulk(self, num: int = DEFAULT_BULK_NUMBER) -> list[dict]:
        """Read bulk XLSX records"""
        chunk = []
        for _n in range(0, num):
            try:
                row = next(self.cursor)
            except StopIteration:
                break
            tmp = list()
            for cell in row:
                tmp.append(str(cell.value))
            result = dict(zip(self.keys, tmp, strict=False))
            chunk.append(result)
            self.pos += 1
        return chunk

    def _buffer_row(self, record: dict) -> None:
        """Buffer one record for the workbook materialized on close()."""
        for key in record:
            if key not in self._write_keys:
                self._write_keys.append(key)
        self._write_rows.append(dict(record))

    def _materialize_workbook(self) -> None:
        """Write the buffered rows to an .xlsx workbook and save it."""
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet()
        if self._write_keys:
            sheet.append(list(self._write_keys))
        for record in self._write_rows:
            row = []
            for key in self._write_keys:
                value = record.get(key)
                if isinstance(value, _SCALAR_WRITE_TYPES) or hasattr(value, "isoformat"):
                    row.append(value)
                else:
                    row.append(str(value))
            sheet.append(row)
        target = self.filename if self.filename is not None else self.fobj
        workbook.save(target)
        workbook.close()

    def write(self, record: dict) -> None:
        """Write single XLSX record (buffered until close)."""
        if self._validation_hooks:
            validated = self._apply_validation_hooks(record)
            if validated is None:  # Skipped
                return
            record = validated
        self._buffer_row(record)

    def write_bulk(self, records: list[dict]) -> None:
        """Write bulk XLSX records (buffered until close)."""
        if records is None:
            raise TypeError("write_bulk() requires a list of records, got None")
        if self._validation_hooks:
            validated_records = []
            for record in records:
                validated = self._apply_validation_hooks(record)
                if validated is not None:  # Not skipped
                    validated_records.append(validated)
            records = validated_records
        for record in records:
            self._buffer_row(record)

    def close(self):
        """Materialize buffered writes, then close the workbook/file."""
        if self.mode in ("w", "wr") and getattr(self, "_write_rows", None) is not None:
            self._materialize_workbook()
            self._write_rows = None
        if getattr(self, "workbook", None) is not None:
            self.workbook.close()
            self.workbook = None
        super().close()
